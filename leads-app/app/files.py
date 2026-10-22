"""Лиды из файлов: Excel (.xlsx, .xls), CSV, Word (.docx), текст (.txt).

Данные берутся как есть, без распознавания. Таблица: ищем строку заголовка (не обязательно
первую) по тем же названиям колонок, что и на скринах, каждая строка ниже — лид.
Нет таблицы с телефоном — разбираем текст как сообщение (ocr.parse_free_text).
"""
import csv
import io
from datetime import date, datetime
from pathlib import Path

from . import ocr

IMAGE_EXT = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".gif", ".tif", ".tiff"}
DOC_EXT = {".xlsx", ".xlsm", ".xls", ".csv", ".docx", ".txt"}
HEADER_SCAN_ROWS = 30  # строка заголовка ищется среди первых строк листа

# Заголовок ячейки может быть длиннее, чем на скрине: «Номер телефона», «Сумма кредита»
CONTAINS = [("phone", ["телефон", "phone"]), ("lead_date", ["дата"]), ("org_form", ["опф", "форма"]),
            ("amount", ["сумм"]), ("name", ["фио", "компан", "наименован", "клиент", "контрагент"])]


def cell_text(v):
    """Значение ячейки -> строка. Даты — ISO, целые числа без «.0» (телефон 79161234567.0)."""
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d %H:%M") if (v.hour or v.minute) else v.strftime("%Y-%m-%d")
    if isinstance(v, date):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def cell_field(text):
    """Заголовок ячейки -> поле базы (как на скринах, плюс поиск слова внутри заголовка)."""
    text = text.strip()
    if not text:
        return None
    f = ocr.header_field(text)
    if f:
        return f
    low = text.lower()
    return next((f for f, keys in CONTAINS if any(k in low for k in keys)), None)


def rows_from_grid(grid):
    """Таблица (список строк ячеек) -> лиды. None — в таблице нет заголовка с телефоном."""
    for h, header in enumerate(grid[:HEADER_SCAN_ROWS]):
        cols = {}
        for i, v in enumerate(header):
            f = cell_field(cell_text(v))
            if f and f not in cols:  # первая подходящая колонка; «Дата рождения» после «Даты лида» не мешает
                cols[f] = i
        if "phone" in cols:
            break
    else:
        return None
    rows = []
    for line in grid[h + 1:]:
        raw = {f: cell_text(line[i]) if i < len(line) else "" for f, i in cols.items()}
        if not any(raw.values()):
            continue  # пустая строка
        if not ocr.norm_phone(raw["phone"]):
            if any(v for k, v in raw.items() if k != "phone"):  # строка с данными, но без телефона — покажем
                rows.append(ocr.build_row(raw))
            continue
        rows.append(ocr.build_row(raw))
    return rows


def grid_text(grid):
    """Таблица без заголовка -> текст: каждая строка таблицы — строка текста."""
    return "\n".join(" ".join(cell_text(v) for v in row if cell_text(v)) for row in grid)


def parse_grids(grids):
    """Несколько таблиц (листы Excel, таблицы Word) -> лиды."""
    rows, loose = [], []
    for g in grids:
        found = rows_from_grid(g)
        if found is None:
            loose.append(grid_text(g))
        else:
            rows += found
    if loose:
        rows += ocr.parse_free_text("\n".join(loose))
    return rows


def read_xlsx(f):
    from openpyxl import load_workbook
    wb = load_workbook(f, data_only=True, read_only=True)
    return [[list(r) for r in ws.iter_rows(values_only=True)] for ws in wb.worksheets]


def read_xls(f):
    import xlrd
    book = xlrd.open_workbook(file_contents=f.read())
    grids = []
    for sh in book.sheets():
        grid = []
        for r in range(sh.nrows):
            row = []
            for c in range(sh.ncols):
                cell = sh.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    row.append(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode))
                else:
                    row.append(cell.value)
            grid.append(row)
        grids.append(grid)
    return grids


def decode(data):
    for enc in ("utf-8-sig", "cp1251"):  # Excel в России сохраняет CSV в cp1251
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            pass
    return data.decode("utf-8", "replace")


def read_csv(f):
    text = decode(f.read())
    try:
        delimiter = csv.Sniffer().sniff(text[:4096], delimiters=";,\t").delimiter
    except csv.Error:
        delimiter = ";" if text.count(";") > text.count(",") else ","
    return [list(csv.reader(io.StringIO(text), delimiter=delimiter))]


def read_docx(f):
    """Word: таблицы — как таблицы, абзацы — как текст сообщения."""
    import docx
    d = docx.Document(f)
    grids = [[[c.text for c in row.cells] for row in t.rows] for t in d.tables]
    text = "\n".join(p.text for p in d.paragraphs if p.text.strip())
    return grids, text


def parse_document(filename, f):
    """Главная функция: файл -> список лидов (как ocr.parse_image)."""
    ext = Path(filename or "").suffix.lower()
    if ext in (".xlsx", ".xlsm"):
        return parse_grids(read_xlsx(f))
    if ext == ".xls":
        return parse_grids(read_xls(f))
    if ext == ".csv":
        return parse_grids(read_csv(f))
    if ext == ".docx":
        grids, text = read_docx(f)
        rows = parse_grids(grids) if grids else []
        known = {r["phone"] for r in rows}
        # Телефоны из текста вне таблиц — если их ещё нет среди строк таблиц
        rows += [r for r in ocr.parse_free_text(text) if r["phone"] not in known]
        return rows
    if ext == ".txt":
        return ocr.parse_free_text(decode(f.read()))
    if ext == ".doc":
        raise ValueError("старый формат .doc не поддерживается — откройте в Word и сохраните как .docx")
    raise ValueError(f"формат {ext or 'без расширения'} не поддерживается")
