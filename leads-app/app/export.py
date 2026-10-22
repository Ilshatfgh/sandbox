"""Excel: выгрузка базы и загрузка отчёта по отработке.

Один формат на оба направления: выгрузили -> оператор заполнил «Статус», «Перезвонить»,
«Комментарий» -> загрузили обратно. Серые колонки — справочные, их при загрузке не меняем.
"""
from datetime import date, datetime
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .db import STATUSES

COLUMNS = [  # (поле, заголовок, ширина, редактируемая)
    ("phone", "Телефон", 16, False), ("name", "Компания / ФИО", 36, False),
    ("org_form", "Форма", 8, False), ("amount", "Сумма запроса", 15, False),
    ("lead_date", "Дата лида", 17, False),
    ("status", "Статус", 26, True), ("callback_at", "Перезвонить", 17, True),
    ("comment", "Комментарий", 50, True),
]
GREY = PatternFill("solid", fgColor="EEEEEE")
HEAD = PatternFill("solid", fgColor="D9E2F3")

# Заголовки, которые узнаём при загрузке (в т.ч. из чужих отчётов). Сравнение по началу, без регистра.
HEADER_ALIASES = {
    "phone": ["телефон", "phone", "lead_phone", "номер"],
    "name": ["компания", "фио", "клиент", "имя", "наименование"],
    "status": ["статус", "результат", "итог"],
    "callback_at": ["перезвонить", "дата перезвона", "перезвон", "callback"],
    "comment": ["комментарий", "коммент", "примечание"],
}

# Свободные формулировки статуса -> статус из справочника
STATUS_ALIASES = {
    "новый": "Новый", "в работе": "В работе", "работа": "В работе",
    "перезвон": "Перезвонить", "успе": "Успешно", "сделка": "Успешно", "одобр": "Успешно",
    "отказ": "Отказ клиента", "не проход": "Не проходит по критериям",
    "не подход": "Не проходит по критериям", "не соответ": "Не проходит по критериям",
}


def to_xlsx(leads) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Лиды"
    ws.append([c[1] for c in COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = HEAD
    for lead in leads:
        row = []
        for f, *_ in COLUMNS:
            v = lead.get(f)
            if f == "callback_at" and v:  # дата как дата, чтобы Excel давал календарь/сортировку
                try:
                    v = datetime.fromisoformat(v)
                except ValueError:
                    pass
            row.append(v)
        ws.append(row)
    n = ws.max_row
    for i, (_, _, width, editable) in enumerate(COLUMNS, 1):
        col = get_column_letter(i)
        ws.column_dimensions[col].width = width
        if not editable:
            for cell in ws[col][1:]:
                cell.fill = GREY
    for cell in ws["D"][1:]:
        cell.number_format = "# ##0"
    for cell in ws["G"][1:]:
        cell.number_format = "DD.MM.YYYY HH:MM"
    # Выпадающий список статусов (с запасом строк для дописанных вручную)
    dv = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=True)
    dv.error, dv.errorTitle = "Выберите статус из списка", "Неверный статус"
    ws.add_data_validation(dv)
    dv.add(f"F2:F{max(n, 2) + 500}")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{n}"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def norm_status(text):
    """«успешно», «Отказ», «не проходит» -> статус из справочника; None если не узнали."""
    if text is None or str(text).strip() == "":
        return None
    t = str(text).strip().lower()
    for s in STATUSES:
        if t == s.lower():
            return s
    for key, s in STATUS_ALIASES.items():
        if t.startswith(key) or key in t:
            return s
    return None


def norm_datetime(v):
    """Ячейка «Перезвонить» -> «ГГГГ-ММ-ДД ЧЧ:ММ». Пусто -> None; не разобрали -> исходный текст."""
    if v is None or str(v).strip() == "":
        return None
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d %H:%M")
    if isinstance(v, date):
        return v.strftime("%Y-%m-%d") + " 10:00"
    t = str(v).strip()
    for fmt in ("%d.%m.%Y %H:%M", "%d.%m.%Y", "%Y-%m-%d %H:%M", "%Y-%m-%d", "%d.%m.%y", "%d.%m"):
        try:
            dt = datetime.strptime(t, fmt)
            if fmt == "%d.%m":
                dt = dt.replace(year=date.today().year)
            return dt.strftime("%Y-%m-%d %H:%M") if "%H" in fmt else dt.strftime("%Y-%m-%d") + " 10:00"
        except ValueError:
            pass
    return t


def read_report(file):
    """Читаем отчёт: ищем строку заголовков (в первых 10 строках), возвращаем строки как dict."""
    wb = load_workbook(file, data_only=True, read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    for h_idx, header in enumerate(rows[:10]):
        cols = {}
        for i, h in enumerate(header):
            h = str(h or "").strip().lower()
            for field, aliases in HEADER_ALIASES.items():
                if field not in cols and any(h.startswith(a) for a in aliases):
                    cols[field] = i
                    break
        if "status" in cols and ("phone" in cols or "name" in cols):
            break
    else:
        raise ValueError("Не найдены колонки «Телефон» (или «Компания / ФИО») и «Статус»")
    out = []
    for n, r in enumerate(rows[h_idx + 1:], start=h_idx + 2):
        item = {f: (r[i] if i < len(r) else None) for f, i in cols.items()}
        if all(v is None or str(v).strip() == "" for v in item.values()):
            continue  # пустая строка
        item["row_no"] = n
        out.append(item)
    return out
