"""Тесты: нормализация, распознавание тестового скрина, API импорта/статусов/экспорта.
Запуск: python -m pytest tests"""
import io
import sqlite3
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import db, ocr

FIXTURE = Path(__file__).parent / "fixtures" / "crm_table.png"
FIXTURE_REAL = Path(__file__).parent / "fixtures" / "crm_real.png"  # раскладка реального скрина


@pytest.mark.parametrize("raw,exp", [
    ("79161234567", "+79161234567"), ("8 (916) 123-45-67", "+79161234567"),
    ("+7 916 123 45 67", "+79161234567"), ("9161234567", "+79161234567"), ("12345", None),
])
def test_norm_phone(raw, exp):
    assert ocr.norm_phone(raw) == exp


def test_norm_other():
    assert ocr.norm_amount("1 200 000,00 ₽") == 1200000
    assert ocr.norm_org_form("000") == "ООО"
    assert ocr.norm_org_form("ип") == "ИП"
    assert ocr.norm_date("2026-10-2014:32") == "2026-10-20 14:32"
    assert ocr.norm_date("20.10.2026") == "2026-10-20"
    assert ocr.strip_org_form("ООО \"Ромашка\"", "ООО") == "Ромашка"
    assert ocr.norm_org_form("un") == "ИП"           # мелкое «ИП» OCR читает как «un»
    assert ocr.norm_org_form("null") is None
    assert ocr.guess_org_form("U", None) == (None, None)  # серый «null» — не ошибка
    assert ocr.guess_org_form("wn", "ИП Петрова Анна") == ("ИП", None)
    assert ocr.guess_org_form("wn", "wn Сидоров Олег Игоревич") == ("ИП", None)  # мусор вместо «ИП»
    assert ocr.guess_org_form(None, "Сидоров Олег Игоревич") == (None, None)      # просто ФИО — не ИП
    assert ocr.norm_phone("7. 79991697749") == "+79991697749"  # уголок Excel перед номером
    assert ocr.norm_date("05 07 2026 15:48:05") == "2026-07-05 15:48"
    assert ocr.norm_date("12.янв") == "2026-01-12" or ocr.norm_date("12.янв").endswith("-01-12")
    assert ocr.norm_date("3 марта 2025") == "2025-03-03"


def test_parse_table_screenshot():
    rows = ocr.parse_image(FIXTURE)
    got = [(r["lead_date"], r["phone"], r["name"], r["org_form"], r["amount"]) for r in rows]
    assert got == [
        ("2026-10-20 14:32", "+79161234567", "Иванов Иван Иванович", "ИП", 500000),
        ("2026-10-20 15:01", "+79035550011", "Ромашка", "ООО", 1200000),
        ("2026-10-21 09:15", "+79269998877", "Вектор Плюс", "ООО", 3500000),
    ]
    assert [r["inn"] for r in rows] == ["771234567890", "7701234567", "7809876543"]


def test_parse_real_layout():
    rows = ocr.parse_image(FIXTURE_REAL)
    got = [(r["lead_date"], r["phone"], r["name"], r["org_form"], r["amount"]) for r in rows]
    assert got == [
        ("2026-10-18 23:41", "+79001112233", None, None, 1000000),
        ("2026-10-18 23:36", "+79002223344", "Петрова Анна Сергеевна", "ИП", 1500000),
        ("2026-10-18 23:05", "+79003334455", "Сидоров Олег Игоревич", "ИП", 100000),
    ]
    assert [r["inn"] for r in rows] == [None, "110402700000", "971500000000"]


def test_free_text_fallback():
    rows = ocr.parse_free_text("Новая заявка 21.10.2026\nООО Альфа, тел 8 (999) 111-22-33, сумма 750 000 руб")
    assert len(rows) == 1
    r = rows[0]
    assert (r["phone"], r["org_form"], r["amount"], r["lead_date"]) == ("+79991112233", "ООО", 750000, "2026-10-21")
    # ИНН: с подписью и без; 10 цифр, начинающиеся с 9, — это телефон, не ИНН
    assert ocr.parse_free_text("ИП Петров, ИНН: 771234567890, тел. 89991112233")[0]["inn"] == "771234567890"
    assert ocr.parse_free_text("ООО Луч 9991112233 7701234567")[0]["inn"] == "7701234567"
    assert ocr.parse_free_text("ООО Луч 9991112233")[0]["inn"] is None


def test_inn():
    assert ocr.norm_inn("7701 234 567") == "7701234567"
    assert ocr.norm_inn(771234567890) == "771234567890"
    assert ocr.norm_inn("12345") is None
    r = ocr.build_row({"phone": "79001112233", "inn": "123"})
    assert r["inn"] is None and any("ИНН не распознан" in w for w in r["warnings"])
    # ИНН, прилипший к названию на скрине, не теряется
    r = ocr.build_row({"phone": "79001112233", "name": "771234567890ИП Иванов Иван Иванович"})
    assert (r["inn"], r["name"], r["org_form"]) == ("771234567890", "Иванов Иван Иванович", "ИП")


@pytest.fixture
def client(tmp_path, monkeypatch):
    from app import main
    monkeypatch.setattr(main, "con", db.connect(tmp_path / "t.db"))
    return TestClient(main.app)


def test_flow(client):
    # 1. Распознать скрин
    with open(FIXTURE, "rb") as f:
        res = client.post("/api/parse", files=[("files", ("a.png", f, "image/png"))]).json()
    assert len(res["rows"]) == 3 and not res["rows"][0]["dup_lead"]

    # 2. Сохранить
    rows = [{k: r[k] for k in ("lead_date", "phone", "name", "org_form", "amount")} for r in res["rows"]]
    assert client.post("/api/leads/import", json=rows).json()["added"] == 3

    # 3. Повторная загрузка — дубли найдены и пропущены
    with open(FIXTURE, "rb") as f:
        res2 = client.post("/api/parse", files=[("files", ("a.png", f, "image/png"))]).json()
    assert all(r["dup_lead"] for r in res2["rows"])
    assert client.post("/api/leads/import", json=rows).json() == {"added": 0, "updated": 0, "skipped": 3, "errors": []}

    # 4. Статусы: «Перезвонить» без даты нельзя, с датой можно
    lead = client.get("/api/leads").json()[0]
    assert client.patch(f"/api/leads/{lead['id']}", json={"status": "Перезвонить"}).status_code == 400
    r = client.patch(f"/api/leads/{lead['id']}", json={"status": "Перезвонить", "callback_at": "2020-01-01 10:00"})
    assert r.status_code == 200
    assert [l["id"] for l in client.get("/api/leads?callback_due=true").json()] == [lead["id"]]
    assert client.patch(f"/api/leads/{lead['id']}", json={"status": "Бред"}).status_code == 400

    # 5. Экспорт в Excel
    wb = load_workbook(io.BytesIO(client.get("/api/export.xlsx").content))
    ws = wb.active
    assert ws.max_row == 4 and ws["A1"].value == "Телефон" and ws["C1"].value == "ИНН" and ws["G1"].value == "Статус"


def test_inn_api(client):
    """ИНН: сохранение, правка, поиск, экспорт; база старой версии получает колонку ИНН."""
    client.post("/api/leads/import", json=[{"phone": "79001112233", "name": "Альфа", "inn": "7701 234 567"}])
    lead = client.get("/api/leads").json()[0]
    assert lead["inn"] == "7701234567"
    assert [l["id"] for l in client.get("/api/leads?q=770123").json()] == [lead["id"]]
    assert client.patch(f"/api/leads/{lead['id']}", json={"inn": "12"}).status_code == 400
    assert client.patch(f"/api/leads/{lead['id']}", json={"inn": "771234567890"}).json()["inn"] == "771234567890"
    # Обновление дубля из загрузки дописывает ИНН
    client.post("/api/leads/import", json=[{"phone": "79001112233", "inn": "7809876543", "action": "update"}])
    assert client.get("/api/leads").json()[0]["inn"] == "7809876543"
    ws = load_workbook(io.BytesIO(client.get("/api/export.xlsx").content)).active
    assert (ws["C1"].value, ws["C2"].value) == ("ИНН", "7809876543")


def test_migrate_old_db(tmp_path):
    path = tmp_path / "old.db"
    old = sqlite3.connect(path)
    old.executescript(db.SCHEMA.replace("    inn         TEXT,               -- ИНН: 10 цифр (юрлицо) или 12 (ИП / физлицо)\n", ""))
    old.execute("INSERT INTO leads (phone, name, created_at, updated_at) VALUES ('+79001112233', 'Старый', '', '')")
    old.commit(); old.close()
    con = db.connect(path)
    assert con.execute("SELECT name, inn FROM leads").fetchone()[:] == ("Старый", None)
    db.insert_lead(con, {"phone": "+79002223344", "inn": "7701234567"})


def test_report_roundtrip(client):
    """Экспорт -> оператор заполняет статусы в Excel -> загрузка -> предпросмотр -> применение."""
    rows = [{"phone": p, "name": n} for p, n in
            [("79001112233", "Альфа"), ("79002223344", "Бета"), ("79003334455", "Гамма"), ("79004445566", "Дельта")]]
    client.post("/api/leads/import", json=rows)
    wb = load_workbook(io.BytesIO(client.get("/api/export.xlsx").content))
    ws = wb.active
    by_phone = {ws.cell(r, 1).value: r for r in range(2, ws.max_row + 1)}
    ws.cell(by_phone["+79001112233"], 7).value = "Успешно"
    ws.cell(by_phone["+79002223344"], 7).value = "отказ"            # свободная формулировка
    ws.cell(by_phone["+79002223344"], 9).value = "Дорого"
    ws.cell(by_phone["+79003334455"], 7).value = "Перезвонить"      # без даты — ошибка
    ws.cell(by_phone["+79004445566"], 7).value = "Перезвонить"
    ws.cell(by_phone["+79004445566"], 8).value = "25.10.2026 15:00"
    ws.append(["+79990000000", "Нет в базе", None, None, None, None, "Успешно"])
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)

    res = client.post("/api/report/parse", files={"file": ("r.xlsx", buf, "application/octet-stream")}).json()["rows"]
    by = {(r["lead"] or {}).get("phone") or r["raw_phone"]: r for r in res}
    assert by["+79001112233"]["changes"] == {"status": "Успешно"}
    assert by["+79002223344"]["changes"] == {"status": "Отказ клиента", "comment": "Дорого"}
    assert by["+79003334455"]["errors"] == ["Статус «Перезвонить» без даты"]
    assert by["+79004445566"]["changes"] == {"status": "Перезвонить", "callback_at": "2026-10-25 15:00"}
    assert by["+79990000000"]["errors"] == ["Клиент не найден в базе"]

    ok = [r for r in res if r["changes"] and not r["errors"]]
    assert client.post("/api/report/apply", json=[{"id": r["lead"]["id"], **r["changes"]} for r in ok]).json() == {"updated": 3}
    st = {l["phone"]: (l["status"], l["comment"]) for l in client.get("/api/leads").json()}
    assert st["+79002223344"] == ("Отказ клиента", "Дорого") and st["+79003334455"][0] == "Новый"


def test_backup(tmp_path, monkeypatch):
    from app import backup
    src = tmp_path / "leads.db"
    con = db.connect(src)
    db.insert_lead(con, {"phone": "+79001112233", "name": "Альфа"})
    con.commit()
    folder = tmp_path / "backups"
    for i in range(35):  # старые копии сверх лимита удаляются
        (folder / f"leads_2020-01-{i:02d}.db").parent.mkdir(exist_ok=True)
        (folder / f"leads_2020-01-{i:02d}.db").write_bytes(b"")
    target = backup.make_backup(src, folder)
    assert len(list(folder.glob("leads_*.db"))) == backup.KEEP
    assert sqlite3.connect(target).execute("SELECT name FROM leads").fetchone()[0] == "Альфа"
    assert backup.last_backup(folder) == target.name


def test_export_selected(client):
    """Экспорт выбранных строк (как их отфильтровал интерфейс), имя файла — по периоду."""
    rows = [{"phone": f"7900111220{i}", "name": f"Лид {i}", "lead_date": f"2026-10-0{i}"} for i in range(1, 6)]
    client.post("/api/leads/import", json=rows)
    ids = {l["name"]: l["id"] for l in client.get("/api/leads").json()}
    r = client.post("/api/export.xlsx", json={"ids": [ids["Лид 3"], ids["Лид 1"]],
                                              "date_from": "2026-10-01", "date_to": "2026-10-03"})
    assert 'filename="leads_2026-10-01_2026-10-03.xlsx"' in r.headers["content-disposition"]
    ws = load_workbook(io.BytesIO(r.content)).active
    assert [ws.cell(i, 2).value for i in (2, 3)] == ["Лид 3", "Лид 1"] and ws.max_row == 3


INNS = {"excel_leads.png": ["410100000001", "710000000002", "4100000003"],
        "excel_fio.png": ["590800000004", "701700000005", "240100000006"],
        "telegram_msg.png": ["6600000007"]}  # в сообщении ИНН без подписи — отдельное 10-значное число


@pytest.mark.parametrize("fixture,expected", [
    # Выгрузка CRM в Excel: заголовки вплотную, зелёные уголки, ИНН рядом с суммой
    ("excel_leads.png", [("2026-07-05 23:53", "+79001234501", "Смирнова Ольга Петровна", "ИП", 1000000),
                         ("2026-07-05 21:50", "+79001234502", "Кузнецов Павел Андреевич", "ИП", 1500000),
                         ("2026-07-05 21:35", "+79001234503", "ЛУЧ", "ООО", 2000000)]),
    # Таблица «Дата | ФИО | Телефон»: заголовки по центру, дата «12.янв», просто ФИО — не ИП
    ("excel_fio.png", [("-01-12", "+79001234504", "Орлова Мария Ивановна", None, None),
                       ("-01-12", "+79001234505", "Степан", None, None),
                       ("-01-13", "+79001234506", "Лебедев Игорь Олегович", None, None)]),
    # Сообщение без таблицы: ИНН из 10 цифр не телефон
    ("telegram_msg.png", [(None, "+79001234507", "Сафина Алия Ринатовна", None, None)]),
])
def test_other_formats(fixture, expected):
    rows = ocr.parse_image(Path(__file__).parent / "fixtures" / fixture)
    got = [(r["lead_date"], r["phone"], r["name"], r["org_form"], r["amount"]) for r in rows]
    assert [r["inn"] for r in rows] == INNS[fixture]
    # дата без года (12.янв) зависит от текущего года — сравниваем окончание
    got = [(g[0][4:] if e[0] and e[0].startswith("-") and g[0] else g[0],) + g[1:] for g, e in zip(got, expected)]
    assert got == expected


def test_ascii_copy(tmp_path, monkeypatch):
    """Tesseract из папки с русскими буквами копируется в ASCII-папку (ProgramData)."""
    src = tmp_path / "Загрузки" / "tesseract"
    (src / "tessdata").mkdir(parents=True)
    (src / "tesseract.exe").write_bytes(b"x" * 10)
    (src / "tessdata" / "rus.traineddata").write_bytes(b"y" * 20)
    ascii_base = tmp_path / "pd"
    monkeypatch.setenv("ProgramData", str(ascii_base))
    target = ocr.ascii_copy(src)
    assert str(target).isascii() and (target / "tessdata" / "rus.traineddata").stat().st_size == 20
    assert ocr.ascii_copy(src) == target  # повторно — та же папка, без лишнего копирования
    assert ocr.ascii_copy(ascii_base) == ascii_base  # путь без русских букв — как есть


FIX = Path(__file__).parent / "fixtures"


def test_documents():
    from app.files import parse_document
    got = lambda f: [(r["lead_date"], r["phone"], r["name"], r["org_form"], r["amount"])
                     for r in parse_document(f, open(FIX / f, "rb"))]
    # Excel: 2 листа, на втором заголовок в 4-й строке, телефоны числами, даты — даты Excel
    assert got("leads.xlsx") == [
        ("2026-07-05 23:53", "+79001110001", "Смирнова Ольга Петровна", "ИП", 1000000),
        ("2026-07-05 21:50", "+79001110002", "ЛУЧ", "ООО", 1500000),
        ("2026-10-01 12:30", "+79001110003", "Орлов Иван Сергеевич", None, 300000),
        ("2026-10-02", "+79001110004", "Белова Анна", None, 750000),
        ("2026-10-03", None, "Без телефона", None, 100000),   # покажется с предупреждением, не сохранится
    ]
    assert [r["inn"] for r in parse_document("leads.xlsx", open(FIX / "leads.xlsx", "rb"))][:2] == \
        ["410100000001", "710000000002"]  # колонка «инн лида»
    # CSV из русского Excel: cp1251, «;», телефоны в разной записи
    assert got("leads.csv") == [
        ("2026-10-12", "+79001110005", "Кузнецов Павел Андреевич", "ИП", 2000000),
        ("2026-10-13", "+79001110006", "Альфа", "ООО", 500000),
    ]
    # Word: таблица + клиент в обычном тексте
    assert got("leads.docx") == [
        (None, "+79001110007", "Лебедев Игорь Олегович", None, 400000),
        (None, "+79001110008", "Новикова Мария", None, 1200000),
        (None, "+79001110009", "Вектор", "ООО", 900000),
    ]


def test_excel_without_header():
    """Строки, скопированные из CRM без заголовка: колонки по содержимому (id лидов — не сумма)."""
    from app.files import parse_document
    rows = parse_document("x.xlsx", open(FIX / "leads_noheader.xlsx", "rb"))
    assert [r["inn"] for r in rows] == [None, "110400000011", "7700000012"]  # ИНН по содержимому
    assert [(r["lead_date"], r["phone"], r["name"], r["org_form"], r["amount"]) for r in rows] == [
        ("2026-10-20 09:48", "+79001110010", None, None, 50000),
        ("2026-10-20 10:05", "+79001110011", "Петров Пётр Петрович", "ИП", 1500000),
        ("2026-10-20 11:00", "+79001110012", "ЗАРЯ", "ООО", 2000000),
    ]
    assert rows[0]["warnings"][0].startswith("Файл без заголовков")


def test_big_excel_upload(client, tmp_path):
    """Файл на 1500 лидов: все разобраны, дубли внутри файла и с базой помечены."""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(["Телефон", "ФИО", "Сумма"])
    for i in range(1500):
        ws.append([f"7900{i:07d}", f"Клиент {i}", 100000 + i])
    ws.append(["79000000001", "Повтор", 1])  # дубль внутри файла
    path = tmp_path / "big.xlsx"
    wb.save(path)
    client.post("/api/leads/import", json=[{"phone": "79000000002", "name": "Уже в базе"}])
    with open(path, "rb") as f:
        res = client.post("/api/parse", files=[("files", ("big.xlsx", f, "application/octet-stream"))]).json()
    rows = res["rows"]
    assert len(rows) == 1501 and not res["errors"]
    assert rows[-1]["dup_in_batch"] and rows[2]["dup_lead"]["name"] == "Уже в базе"
    imp = [{k: r[k] for k in ("lead_date", "phone", "name", "org_form", "amount")} for r in rows[:-1]]
    assert client.post("/api/leads/import", json=imp).json()["added"] == 1499


def test_unsupported_file(client):
    res = client.post("/api/parse", files=[("files", ("old.doc", b"xx", "application/msword"))]).json()
    assert res["rows"] == [] and "сохраните как .docx" in res["errors"][0]
